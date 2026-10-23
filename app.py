# =====================================================================
# app.py - El "cerebro" de nuestra tienda de ascensores (proyecto SENA)
# Aquí está todo lo del servidor: conexión a la base de datos, login,
# el CRUD de productos y usuarios, validaciones y los reportes PDF/Excel.
# =====================================================================

import os                       # Para manejar rutas de carpetas y leer variables del sistema
import re                       # Para las expresiones regulares (lo usamos para revisar correos)
from io import BytesIO          # Una "memoria" donde armamos el PDF y el Excel sin guardarlos en disco
from datetime import datetime   # Para poner la fecha en los reportes

# De Flask traemos todo lo que necesitamos para la página web
from flask import Flask, render_template, request, redirect, session, flash, url_for, send_file, get_flashed_messages
from markupsafe import Markup, escape                            # Para devolver HTML seguro desde Python (viene con Flask)
import mysql.connector                                           # El "puente" entre Python y MySQL
from mysql.connector import Error as MySQLError, IntegrityError  # Los errores que puede botar MySQL
from openpyxl import Workbook                                    # Librería para crear archivos de Excel
from openpyxl.styles import Font                                 # Para poner letra en negrita en el Excel
from fpdf import FPDF                                            # Librería para crear archivos PDF

# ---------------------------------------------------------------------
# Configuración básica de Flask
# ---------------------------------------------------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))  # La carpeta donde está este archivo
TEMPLATE_DIR = os.path.join(BASE_DIR, "templates")     # La carpeta de los HTML
STATIC_DIR = os.path.join(BASE_DIR, "static")          # La carpeta del CSS, JS e imágenes

# Creamos la app de Flask y le decimos dónde están los HTML y los archivos estáticos
app = Flask(__name__, template_folder=TEMPLATE_DIR, static_folder=STATIC_DIR)
# Esta clave la usa Flask para "firmar" la sesión y que nadie la pueda falsificar
app.secret_key = os.environ.get("SECRET_KEY", "clave_secreta_proyecto_sena")  # En la nube Render nos da una clave propia
# OJO: las rutas de la página tienen máximo dos niveles (ej: /producto_editar/5) a propósito.
# Así el link "../static/css/styles.css" de los HTML siempre llega a la carpeta static.
# Además, varias páginas tienen una "ruta espejo" /templates/archivo.html: los links del menú
# apuntan a "../templates/archivo.html", y eso funciona igual abriendo el HTML suelto en
# Visual Studio Code que con Flask corriendo.

# ---------------------------------------------------------------------
# CONEXIÓN A LA BASE DE DATOS (esto es lo que más preguntan en la sustentación)
# Aquí van los datos para que Python encuentre a MySQL.
# os.environ.get("X", "valor") significa: "si existe la variable X en el sistema úsela,
# y si no, use el valor por defecto". Así sirve igual en el PC (con XAMPP) y en la nube.
#
# En la nube NO tocamos este código: los datos de la base de Aiven se escriben en Render,
# en la sección "Environment" (variables de entorno), y Python los lee de ahí:
#   DB_HOST     -> el Host que muestra Aiven (algo como mysql-xxxx.aivencloud.com)
#   DB_PORT     -> el Port que muestra Aiven (un número largo, no es 3306)
#   DB_USER     -> avnadmin
#   DB_PASSWORD -> la contraseña que muestra Aiven
#   DB_NAME     -> defaultdb
# Aiven exige conexión cifrada (SSL); el conector de MySQL la usa solito cuando el servidor la pide.
# ---------------------------------------------------------------------
DB_CONFIG = {
    "host": os.environ.get("DB_HOST", "localhost"),              # Dónde está el servidor MySQL
    "port": int(os.environ.get("DB_PORT", "3306")),              # Puerto de MySQL (3306 es el normal de XAMPP)
    "user": os.environ.get("DB_USER", "root"),                   # Con qué usuario entramos
    "password": os.environ.get("DB_PASSWORD", ""),               # La contraseña de ese usuario
    "database": os.environ.get("DB_NAME", "tienda_ascensores"),  # El nombre de nuestra base
    "charset": "utf8mb4",                                        # Para que guarde bien tildes, ñ y emojis
    "collation": "utf8mb4_general_ci",                           # Esta sirve en MySQL y en MariaDB (la que trae XAMPP)
    "use_pure": True,                                            # Usamos el conector 100% Python: muestra los errores más claritos
}


def obtener_conexion():
    # Abre la conexión con MySQL usando los datos de arriba (los ** "desempacan" el diccionario)
    return mysql.connector.connect(**DB_CONFIG)


def consultar(sql, parametros=(), uno=False):
    # Función ayudante para los SELECT: así no repetimos el mismo código en todo lado
    conexion = obtener_conexion()                # 1. Abrimos la conexión
    cursor = conexion.cursor(dictionary=True)    # 2. Pedimos un cursor que devuelva diccionarios (fila["nombre"])
    cursor.execute(sql, parametros)              # 3. Ejecutamos la consulta (los %s evitan inyección SQL)
    resultado = cursor.fetchone() if uno else cursor.fetchall()  # 4. Traemos una fila o todas
    cursor.close()                               # 5. Cerramos el cursor
    conexion.close()                             # 6. Cerramos la conexión (siempre hay que cerrar la puerta)
    return resultado                             # 7. Devolvemos lo que encontramos


def ejecutar(sql, parametros=()):
    # Función ayudante para INSERT, UPDATE y DELETE (las que cambian datos)
    conexion = obtener_conexion()                # Abrimos la conexión
    cursor = conexion.cursor()                   # Pedimos un cursor normalito
    try:
        cursor.execute(sql, parametros)          # Ejecutamos la orden
        conexion.commit()                        # commit = "guarde los cambios de verdad"
        return cursor.rowcount                   # Devolvemos cuántas filas se cambiaron (0 = no encontró nada)
    finally:
        cursor.close()                           # Pase lo que pase, cerramos el cursor
        conexion.close()                         # Y cerramos la conexión


def crear_tablas_si_faltan():
    # En la nube no tenemos phpMyAdmin para correr el script SQL a mano.
    # Entonces, cuando arranca la página, revisamos si existe la tabla "usuarios";
    # si NO existe (base nuevecita), corremos nosotros mismos el archivo tienda_online.sql.
    # Si ya existe, no hacemos nada (así no se nos devuelven datos que el admin borró).
    try:
        if consultar("SHOW TABLES LIKE 'usuarios'"):            # Si la tabla ya está...
            return                                             # ...no hay nada que hacer
        with open(os.path.join(BASE_DIR, "tienda_online.sql"), encoding="utf8") as archivo:
            script = archivo.read()                            # Leemos todo el script SQL
        script = re.sub(r"--.*", "", script)                   # Le quitamos los comentarios (-- hasta el final de la línea)
        conexion = obtener_conexion()
        cursor = conexion.cursor()
        for orden in script.split(";"):                        # Cada orden SQL termina en punto y coma
            orden = orden.strip()                              # Le quitamos espacios sobrantes
            if not orden or orden.upper().startswith(("CREATE DATABASE", "USE ")):
                continue                                       # Saltamos las vacías y las de crear/usar la base
            cursor.execute(orden)                              # Ejecutamos la orden
        conexion.commit()                                      # Guardamos todo
        cursor.close()
        conexion.close()
        print("Tablas creadas y datos de prueba cargados.")    # Mensajito en la consola
    except Exception as error:                                 # Si la base no responde, no tumbamos la página
        print("No se pudieron revisar/crear las tablas:", error)


crear_tablas_si_faltan()   # La llamamos una vez, apenas arranca la aplicación


# ---------------------------------------------------------------------
# Funciones de control de acceso (quién puede entrar a qué)
# ---------------------------------------------------------------------
def esta_logueado():
    # Si en la sesión está guardado el id del usuario, es porque ya inició sesión
    return "usuario_id" in session


def es_admin():
    # Revisa si el que está logueado tiene el rol de administrador
    return session.get("rol") == "administrador"


# ---------------------------------------------------------------------
# Validaciones del BackEnd (aunque el HTML ya valida, el servidor vuelve a revisar
# porque alguien se puede saltar el HTML fácilito desde el navegador)
# ---------------------------------------------------------------------
PATRON_CORREO = re.compile(r"^[\w\.\-]+@[\w\-]+\.[\w\.\-]+$")    # Molde de un correo: algo@algo.algo
PATRON_TEXTO = re.compile(r"^[A-Za-zÁÉÍÓÚáéíóúÑñÜü ]+$")         # Solo letras y espacios (para nombres)
PATRON_TELEFONO = re.compile(r"^\+?[0-9 ]{7,15}$")                # Números, espacios y un + opcional


def correo_valido(correo):
    # Devuelve True si el correo tiene forma de correo
    return bool(PATRON_CORREO.match(correo))


def validar_datos_usuario(nombre, apellido, telefono):
    # Revisa nombre, apellido y teléfono; devuelve una lista con los errores que encuentre
    errores = []                                                       # Arrancamos sin errores
    if not nombre or not PATRON_TEXTO.match(nombre) or len(nombre) > 100:
        errores.append("El nombre solo puede tener letras y máximo 100 caracteres.")
    if not apellido or not PATRON_TEXTO.match(apellido) or len(apellido) > 100:
        errores.append("El apellido solo puede tener letras y máximo 100 caracteres.")
    if telefono and not PATRON_TELEFONO.match(telefono):               # El teléfono es opcional, solo se revisa si lo escribieron
        errores.append("El teléfono debe tener entre 7 y 15 números.")
    return errores                                                     # Si la lista sale vacía, todo bien


def leer_entero(texto, permitir_vacio=False):
    # Convierte texto a número entero; si no se puede, devuelve None (señal de error)
    if texto == "" and permitir_vacio:   # Si está vacío y se vale que esté vacío...
        return ""                        # ...lo dejamos vacío
    try:
        return int(texto)                # Intentamos volverlo número
    except ValueError:                   # Si escribieron letras, int() explota y caemos aquí
        return None


def validar_producto(formulario):
    # Lee y revisa todos los campos del formulario de productos.
    # Devuelve (errores, datos): la lista de errores y un diccionario con los datos limpios.
    datos = {
        "categoria_id": formulario.get("categoria_id", "").strip(),             # Categoría escogida
        "nombre": formulario.get("nombre", "").strip(),                         # Nombre del producto
        "modelo": formulario.get("modelo", "").strip().upper(),                 # Modelo en MAYÚSCULAS
        "descripcion": formulario.get("descripcion", "").strip(),               # Descripción
        "precio": formulario.get("precio", "").strip(),                         # Precio (llega como texto)
        "stock": formulario.get("stock", "").strip(),                           # Stock (llega como texto)
        "capacidad_carga_kg": formulario.get("capacidad_carga_kg", "").strip(), # Kilos (opcional)
        "capacidad_personas": formulario.get("capacidad_personas", "").strip(), # Personas (opcional)
        "activo": 1 if formulario.get("activo") else 0,                         # Checkbox: si viene marcado es 1
    }
    errores = []                                                                # Lista de errores vacía

    if leer_entero(datos["categoria_id"]) is None:                             # La categoría tiene que ser un número
        errores.append("Debes escoger una categoría.")
    if len(datos["nombre"]) < 3 or len(datos["nombre"]) > 150:                 # Nombre ni muy corto ni muy largo
        errores.append("El nombre debe tener entre 3 y 150 caracteres.")
    if not re.match(r"^[A-Z0-9\-]{2,100}$", datos["modelo"]):                   # Modelo tipo RES-300
        errores.append("El modelo solo puede tener letras, números y guiones (ej: RES-300).")

    try:
        precio = float(datos["precio"])                                         # Intentamos volver el precio número
        if precio <= 0:                                                         # No regalamos nada
            errores.append("El precio debe ser mayor que cero.")
    except ValueError:                                                          # Si escribieron letras
        errores.append("El precio debe ser un número.")

    stock = leer_entero(datos["stock"])                                         # El stock tiene que ser entero
    if stock is None or stock < 0:                                              # Y no puede ser negativo
        errores.append("El stock debe ser un número entero mayor o igual a cero.")

    for campo in ("capacidad_carga_kg", "capacidad_personas"):                  # Los dos campos opcionales de ascensores
        valor = leer_entero(datos[campo], permitir_vacio=True)                  # Pueden ir vacíos
        if valor is None or (valor != "" and valor <= 0):                       # Pero si escriben algo, debe ser > 0
            errores.append("Las capacidades deben ser números enteros mayores que cero (o dejarlas vacías).")
            break                                                               # Con un mensaje basta
        datos[campo] = valor if valor != "" else None                           # Vacío en el form = NULL en la base

    return errores, datos                                                       # Devolvemos las dos cosas


# Filtro para mostrar la plata bonita en los HTML: 25000 -> $ 25.000
@app.template_filter("pesos")
def formato_pesos(valor):
    return "$ " + f"{valor:,.0f}".replace(",", ".")   # Python pone comas de miles, las cambiamos por puntos


# Función para pintar los avisos de flash() en los HTML.
# En los HTML se llama así:  <!-- {{ avisos() }} -->  (metida dentro de un comentario).
# Truco: si hay avisos, devolvemos "-->" para CERRAR el comentario, pintamos los avisos
# y luego "<!--" para volverlo a abrir. Si no hay avisos, no devolvemos nada.
# Así, si abren el HTML suelto en el navegador, no se ve ningún código raro.
@app.template_global()
def avisos():
    mensajes = get_flashed_messages(with_categories=True)   # Sacamos los avisos guardados con flash()
    if not mensajes:                                        # Si no hay ninguno...
        return ""                                           # ...no pintamos nada
    html = "-->"                                            # Cerramos el comentario del HTML
    for categoria, mensaje in mensajes:                     # Un cuadrito por cada aviso
        clase = "mensaje-exito" if categoria == "exito" else "mensaje-error"   # Verde o rojo
        html += f'<div class="{clase}">{escape(mensaje)}</div>'  # escape() evita que metan código HTML en el texto
    return Markup(html + "<!--")                            # Volvemos a abrir el comentario; Markup = "esto es HTML confiable"


# =====================================================================
# PÁGINAS PÚBLICAS (no necesitan login)
# =====================================================================

@app.route("/")                                  # La página principal
@app.route("/templates/index.html")  # Ruta espejo: la usan los links del menú
def index():
    return render_template("index.html")         # Solo mostramos el HTML de inicio


@app.route("/login", methods=["GET", "POST"])    # GET = mostrar el form, POST = cuando le dan "Ingresar"
@app.route("/templates/login.html", methods=["GET", "POST"])  # Ruta espejo: la usan los links del menú
def login():
    if request.method == "GET":                  # Si solo están entrando a la página...
        return render_template("login.html")     # ...les mostramos el formulario

    correo = request.form.get("correo", "").strip().lower()       # Leemos el correo, sin espacios y en minúscula
    contrasena = request.form.get("contrasena", "").strip()       # Leemos la contraseña

    if not correo or not contrasena:                              # Validación backend: nada vacío
        flash("Debes ingresar correo y contraseña.", "error")
        return redirect(url_for("login"))
    if not correo_valido(correo):                                 # Validación backend: correo con forma de correo
        flash("El correo no tiene un formato válido.", "error")
        return redirect(url_for("login"))

    try:
        # Buscamos al usuario por su correo y de una traemos el nombre del rol con un JOIN
        usuario = consultar(
            "SELECT u.id, u.nombre, u.contrasena, u.activo, r.nombre AS rol "
            "FROM usuarios u JOIN roles r ON r.id = u.rol_id "
            "WHERE u.correo = %s",
            (correo,), uno=True)
    except MySQLError:                                            # Si MySQL está apagado o mal configurado
        flash("No se pudo conectar con la base de datos. Intenta más tarde.", "error")
        return redirect(url_for("login"))

    if usuario and usuario["contrasena"] == contrasena:           # Si existe y la contraseña coincide...
        if not usuario["activo"]:                                 # ...pero el admin lo desactivó
            flash("Tu cuenta está desactivada. Habla con el administrador.", "error")
            return redirect(url_for("login"))
        session["usuario_id"] = usuario["id"]                     # Guardamos en la sesión quién es
        session["nombre"] = usuario["nombre"]                     # Su nombre para saludarlo
        session["rol"] = usuario["rol"]                           # Y su rol para saber qué puede ver
        if usuario["rol"] == "administrador":                     # El admin se va para su panel
            return redirect(url_for("admin"))
        return redirect(url_for("productos"))                     # Los demás se van al catálogo

    flash("Correo o contraseña incorrectos.", "error")            # Si no coincidió nada
    return redirect(url_for("login"))


@app.route("/register", methods=["GET", "POST"])                  # Registro de clientes nuevos
@app.route("/templates/register.html", methods=["GET", "POST"])  # Ruta espejo: la usan los links del menú
def register():
    if request.method == "GET":                                   # Si solo entran a ver...
        return render_template("register.html")                   # ...mostramos el formulario

    nombre = request.form.get("nombre", "").strip()               # Leemos todos los campos
    apellido = request.form.get("apellido", "").strip()
    correo = request.form.get("correo", "").strip().lower()
    telefono = request.form.get("telefono", "").strip()
    contrasena = request.form.get("contrasena", "").strip()
    confirmar = request.form.get("confirmar", "").strip()

    errores = validar_datos_usuario(nombre, apellido, telefono)   # Revisamos nombre, apellido y teléfono
    if not correo_valido(correo):                                 # Revisamos el correo
        errores.append("El correo no tiene un formato válido.")
    if len(contrasena) < 6:                                       # Contraseña mínimo de 6
        errores.append("La contraseña debe tener mínimo 6 caracteres.")
    if contrasena != confirmar:                                   # Las dos contraseñas iguales
        errores.append("Las contraseñas no coinciden.")

    if errores:                                                   # Si hubo algún error...
        for error in errores:
            flash(error, "error")                                 # ...los mostramos todos
        return render_template("register.html", datos=request.form)  # Y devolvemos el form con lo que ya escribió

    try:
        # Insertamos al usuario; el rol lo sacamos con un sub-SELECT buscando el rol "cliente"
        ejecutar(
            "INSERT INTO usuarios (nombre, apellido, correo, telefono, contrasena, rol_id) "
            "VALUES (%s, %s, %s, %s, %s, (SELECT id FROM roles WHERE nombre = 'cliente'))",
            (nombre, apellido, correo, telefono or None, contrasena))
    except IntegrityError:                                        # Si el correo ya existe (es UNIQUE)
        flash("Ya existe una cuenta registrada con ese correo.", "error")
        return render_template("register.html", datos=request.form)
    except MySQLError:                                            # Cualquier otro problema con MySQL
        flash("No se pudo completar el registro. Intenta más tarde.", "error")
        return render_template("register.html", datos=request.form)

    flash("Tu cuenta fue creada. Ya puedes iniciar sesión.", "exito")  # Todo salió bien
    return redirect(url_for("login"))


@app.route("/contacto", methods=["GET", "POST"])                  # Formulario de contacto
@app.route("/templates/contacto.html", methods=["GET", "POST"])  # Ruta espejo: la usan los links del menú
def contacto():
    if request.method == "GET":                                   # Solo mostrar
        return render_template("contacto.html")

    nombre = request.form.get("nombre", "").strip()               # Leemos los campos
    correo = request.form.get("correo", "").strip().lower()
    asunto = request.form.get("asunto", "").strip()
    mensaje = request.form.get("mensaje", "").strip()

    errores = []                                                  # Validaciones del backend
    if len(nombre) < 3 or len(nombre) > 100:
        errores.append("El nombre debe tener entre 3 y 100 caracteres.")
    if not correo_valido(correo):
        errores.append("El correo no tiene un formato válido.")
    if len(asunto) > 150:
        errores.append("El asunto puede tener máximo 150 caracteres.")
    if len(mensaje) < 10:
        errores.append("El mensaje debe tener mínimo 10 caracteres.")

    if errores:                                                   # Si algo falló, avisamos
        for error in errores:
            flash(error, "error")
        return render_template("contacto.html", datos=request.form)

    try:
        ejecutar(                                                 # Guardamos el mensaje en la base
            "INSERT INTO mensajes_contacto (nombre, correo, asunto, mensaje) VALUES (%s, %s, %s, %s)",
            (nombre, correo, asunto or None, mensaje))
    except MySQLError:
        flash("No se pudo enviar el mensaje. Intenta más tarde.", "error")
        return render_template("contacto.html", datos=request.form)

    flash("Tu mensaje fue enviado correctamente. Pronto te contactamos.", "exito")
    return redirect(url_for("contacto"))


@app.route("/logout")                         # Cerrar sesión
def logout():
    session.clear()                           # Borramos todo lo que había en la sesión
    return redirect(url_for("index"))         # Y lo mandamos al inicio


# =====================================================================
# CATÁLOGO (necesita login)
# =====================================================================

@app.route("/productos")
@app.route("/templates/productos.html")  # Ruta espejo: la usan los links del menú
def productos():
    if not esta_logueado():                   # Si no ha iniciado sesión, pa' fuera
        flash("Inicia sesión para ver el catálogo.", "error")
        return redirect(url_for("login"))
    try:
        # READ del CRUD: traemos los productos activos con el nombre de su categoría
        lista = consultar(
            "SELECT p.*, c.nombre AS categoria FROM productos p "
            "JOIN categorias c ON c.id = p.categoria_id "
            "WHERE p.activo = TRUE ORDER BY p.nombre")
    except MySQLError:
        flash("No se pudieron cargar los productos.", "error")
        lista = []                            # Si falla, mostramos la página vacía en vez de explotar
    return render_template("productos.html", productos=lista)


# =====================================================================
# PANEL DEL ADMINISTRADOR (solo entra el rol administrador)
# =====================================================================

@app.route("/admin")
@app.route("/templates/admin.html")  # Ruta espejo: la usan los links del menú
def admin():
    if not es_admin():                        # Control de acceso: si no es admin, lo sacamos
        return redirect(url_for("login"))
    try:
        # Contamos cuántos registros hay para mostrar unos numeritos en el panel
        resumen = consultar(
            "SELECT (SELECT COUNT(*) FROM productos) AS productos, "
            "(SELECT COUNT(*) FROM usuarios) AS usuarios, "
            "(SELECT COUNT(*) FROM mensajes_contacto WHERE estado = 'nuevo') AS mensajes_nuevos",
            uno=True)
    except MySQLError:
        resumen = None                        # Si falla, el panel sale sin numeritos
    return render_template("admin.html", resumen=resumen)


# ------------------------- CRUD DE PRODUCTOS -------------------------

@app.route("/admin_productos")                # R de CRUD (Read): listar todos los productos
@app.route("/templates/admin_productos.html")  # Ruta espejo: la usan los links del menú
def admin_productos():
    if not es_admin():
        return redirect(url_for("login"))
    lista = consultar(
        "SELECT p.*, c.nombre AS categoria FROM productos p "
        "JOIN categorias c ON c.id = p.categoria_id ORDER BY p.id")
    return render_template("admin_productos.html", productos=lista)


@app.route("/producto_nuevo", methods=["GET", "POST"])   # C de CRUD (Create): crear producto
@app.route("/templates/producto_form.html", methods=["GET", "POST"])  # Ruta espejo: la usan los links del menú
def producto_nuevo():
    if not es_admin():
        return redirect(url_for("login"))
    categorias = consultar("SELECT id, nombre FROM categorias ORDER BY nombre")  # Para llenar la lista desplegable

    if request.method == "GET":                                  # Mostramos el formulario vacío
        return render_template("producto_form.html", categorias=categorias, producto={"activo": 1}, titulo="Nuevo producto")

    errores, datos = validar_producto(request.form)              # Validamos lo que mandaron
    if errores:                                                  # Si hay errores, volvemos al form con los datos
        for error in errores:
            flash(error, "error")
        return render_template("producto_form.html", categorias=categorias, producto=datos, titulo="Nuevo producto")

    try:
        ejecutar(                                                # INSERT del producto nuevo
            "INSERT INTO productos (categoria_id, nombre, modelo, descripcion, precio, stock, "
            "capacidad_carga_kg, capacidad_personas, activo) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)",
            (datos["categoria_id"], datos["nombre"], datos["modelo"], datos["descripcion"], datos["precio"],
             datos["stock"], datos["capacidad_carga_kg"], datos["capacidad_personas"], datos["activo"]))
    except IntegrityError:                                       # El modelo es UNIQUE, no se puede repetir
        flash("Ya existe un producto con ese modelo.", "error")
        return render_template("producto_form.html", categorias=categorias, producto=datos, titulo="Nuevo producto")

    flash("Producto creado correctamente.", "exito")
    return redirect(url_for("admin_productos"))


@app.route("/producto_editar/<int:id>", methods=["GET", "POST"])  # U de CRUD (Update): editar
def producto_editar(id):
    if not es_admin():
        return redirect(url_for("login"))
    categorias = consultar("SELECT id, nombre FROM categorias ORDER BY nombre")

    if request.method == "GET":                                  # Buscamos el producto para pintar el form lleno
        producto = consultar("SELECT * FROM productos WHERE id = %s", (id,), uno=True)
        if not producto:                                         # Si ese id no existe
            flash("Ese producto no existe.", "error")
            return redirect(url_for("admin_productos"))
        return render_template("producto_form.html", categorias=categorias, producto=producto, titulo="Editar producto")

    errores, datos = validar_producto(request.form)              # Validamos igual que al crear
    datos["id"] = id                                             # Guardamos el id para que el form siga apuntando a este producto
    if errores:
        for error in errores:
            flash(error, "error")
        return render_template("producto_form.html", categorias=categorias, producto=datos, titulo="Editar producto")

    try:
        ejecutar(                                                # UPDATE con los datos nuevos
            "UPDATE productos SET categoria_id=%s, nombre=%s, modelo=%s, descripcion=%s, precio=%s, stock=%s, "
            "capacidad_carga_kg=%s, capacidad_personas=%s, activo=%s WHERE id=%s",
            (datos["categoria_id"], datos["nombre"], datos["modelo"], datos["descripcion"], datos["precio"],
             datos["stock"], datos["capacidad_carga_kg"], datos["capacidad_personas"], datos["activo"], id))
    except IntegrityError:
        flash("Ya existe otro producto con ese modelo.", "error")
        return render_template("producto_form.html", categorias=categorias, producto=datos, titulo="Editar producto")

    flash("Producto actualizado correctamente.", "exito")
    return redirect(url_for("admin_productos"))


@app.route("/producto_eliminar/<int:id>", methods=["POST"])  # D de CRUD (Delete): borrar
def producto_eliminar(id):
    if not es_admin():
        return redirect(url_for("login"))
    try:
        borrados = ejecutar("DELETE FROM productos WHERE id = %s", (id,))  # Borramos el producto
        if borrados:                                                    # Si borró al menos uno
            flash("Producto eliminado.", "exito")
        else:                                                           # Si ese id no existía
            flash("Ese producto no existe.", "error")
    except IntegrityError:
        # Si el producto ya está en un pedido o contrato, la llave foránea no deja borrarlo
        flash("No se puede eliminar porque tiene pedidos o contratos. Mejor desactívalo.", "error")
    return redirect(url_for("admin_productos"))


# ------------------------- GESTIÓN DE USUARIOS -------------------------

@app.route("/admin_usuarios")                    # Listar usuarios
@app.route("/templates/admin_usuarios.html")  # Ruta espejo: la usan los links del menú
def admin_usuarios():
    if not es_admin():
        return redirect(url_for("login"))
    lista = consultar(
        "SELECT u.id, u.nombre, u.apellido, u.correo, u.telefono, u.activo, u.fecha_registro, r.nombre AS rol "
        "FROM usuarios u JOIN roles r ON r.id = u.rol_id ORDER BY u.id")
    return render_template("admin_usuarios.html", usuarios=lista)


@app.route("/usuario_editar/<int:id>", methods=["GET", "POST"])  # Editar un usuario
def usuario_editar(id):
    if not es_admin():
        return redirect(url_for("login"))
    roles = consultar("SELECT id, nombre FROM roles ORDER BY id")       # Para la lista desplegable de roles
    usuario = consultar("SELECT * FROM usuarios WHERE id = %s", (id,), uno=True)
    if not usuario:                                                     # Si no existe, lo devolvemos
        flash("Ese usuario no existe.", "error")
        return redirect(url_for("admin_usuarios"))

    if request.method == "GET":                                         # Mostramos el form con sus datos
        return render_template("usuario_form.html", usuario=usuario, roles=roles)

    nombre = request.form.get("nombre", "").strip()                     # Leemos lo que cambiaron
    apellido = request.form.get("apellido", "").strip()
    telefono = request.form.get("telefono", "").strip()
    rol_id = leer_entero(request.form.get("rol_id", ""))
    activo = 1 if request.form.get("activo") else 0

    errores = validar_datos_usuario(nombre, apellido, telefono)         # Validamos
    if rol_id is None or rol_id not in [r["id"] for r in roles]:        # El rol tiene que ser uno de los que existen
        errores.append("Debes escoger un rol válido.")
    if id == session["usuario_id"] and not activo:                      # El admin no se puede bloquear a sí mismo
        errores.append("No te puedes desactivar a ti mismo.")
    if id == session["usuario_id"] and rol_id != usuario["rol_id"]:     # Ni quitarse el rol de admin (se quedaría por fuera)
        errores.append("No te puedes cambiar el rol a ti mismo.")

    if errores:
        for error in errores:
            flash(error, "error")
        usuario.update(nombre=nombre, apellido=apellido, telefono=telefono, rol_id=rol_id, activo=activo)
        return render_template("usuario_form.html", usuario=usuario, roles=roles)

    ejecutar(                                                           # UPDATE del usuario
        "UPDATE usuarios SET nombre=%s, apellido=%s, telefono=%s, rol_id=%s, activo=%s WHERE id=%s",
        (nombre, apellido, telefono or None, rol_id, activo, id))
    flash("Usuario actualizado correctamente.", "exito")
    return redirect(url_for("admin_usuarios"))


@app.route("/usuario_estado/<int:id>", methods=["POST"])  # Botón rápido para activar/desactivar
def usuario_estado(id):
    if not es_admin():
        return redirect(url_for("login"))
    if id == session["usuario_id"]:                              # No se vale desactivarse uno mismo
        flash("No te puedes desactivar a ti mismo.", "error")
    else:
        cambiados = ejecutar("UPDATE usuarios SET activo = NOT activo WHERE id = %s", (id,))  # NOT activo = le da la vuelta
        if cambiados:
            flash("Se cambió el estado del usuario.", "exito")
        else:
            flash("Ese usuario no existe.", "error")
    return redirect(url_for("admin_usuarios"))


# ------------------------- MENSAJES DE CONTACTO -------------------------

ESTADOS_MENSAJE = ["nuevo", "en_proceso", "resuelto"]            # Los estados que deja la tabla (el ENUM)


@app.route("/admin_mensajes")                                    # Ver los mensajes que llegaron
@app.route("/templates/admin_mensajes.html")  # Ruta espejo: la usan los links del menú
def admin_mensajes():
    if not es_admin():
        return redirect(url_for("login"))
    lista = consultar("SELECT * FROM mensajes_contacto ORDER BY fecha_creacion DESC")  # Los más nuevos primero
    return render_template("admin_mensajes.html", mensajes=lista, estados=ESTADOS_MENSAJE)


@app.route("/mensaje_estado/<int:id>", methods=["POST"])  # Cambiar el estado de un mensaje
def mensaje_estado(id):
    if not es_admin():
        return redirect(url_for("login"))
    estado = request.form.get("estado", "")
    if estado not in ESTADOS_MENSAJE:                            # Validamos que sea un estado permitido
        flash("Estado no válido.", "error")
    else:
        consultar_existe = consultar("SELECT id FROM mensajes_contacto WHERE id = %s", (id,), uno=True)
        if not consultar_existe:                                 # Si ese mensaje no existe
            flash("Ese mensaje no existe.", "error")
        else:
            ejecutar("UPDATE mensajes_contacto SET estado = %s WHERE id = %s", (estado, id))
            flash("Estado del mensaje actualizado.", "exito")
    return redirect(url_for("admin_mensajes"))


# =====================================================================
# REPORTES (Excel y PDF)
# =====================================================================

@app.route("/reporte_productos_excel")                     # Reporte de inventario en Excel
def reporte_productos_excel():
    if not es_admin():
        return redirect(url_for("login"))
    lista = consultar(
        "SELECT p.id, p.modelo, p.nombre, c.nombre AS categoria, p.precio, p.stock, p.activo "
        "FROM productos p JOIN categorias c ON c.id = p.categoria_id ORDER BY p.id")

    libro = Workbook()                                           # Creamos un libro de Excel nuevo
    hoja = libro.active                                          # Tomamos la primera hoja
    hoja.title = "Inventario"                                    # Le ponemos nombre a la hoja
    encabezados = ["ID", "Modelo", "Nombre", "Categoría", "Precio", "Stock", "Valor en bodega", "Activo"]
    hoja.append(encabezados)                                     # La primera fila son los títulos
    for celda in hoja[1]:                                        # Recorremos las celdas de la fila 1...
        celda.font = Font(bold=True)                             # ...y las ponemos en negrita

    for p in lista:                                              # Una fila por cada producto
        hoja.append([p["id"], p["modelo"], p["nombre"], p["categoria"], float(p["precio"]), p["stock"],
                     float(p["precio"]) * p["stock"], "Sí" if p["activo"] else "No"])

    anchos = [6, 14, 38, 28, 14, 8, 18, 8]                       # Ancho de cada columna para que se lea bien
    for indice, ancho in enumerate(anchos):
        hoja.column_dimensions[chr(65 + indice)].width = ancho  # chr(65) es "A", chr(66) es "B", etc.

    archivo = BytesIO()                                          # Armamos el archivo en memoria
    libro.save(archivo)                                          # Guardamos el Excel ahí
    archivo.seek(0)                                              # Rebobinamos al inicio para poder enviarlo
    nombre = f"inventario_{datetime.now():%Y%m%d}.xlsx"          # Nombre con la fecha de hoy
    return send_file(archivo, as_attachment=True, download_name=nombre,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def texto_pdf(valor):
    # Las fuentes básicas del PDF no aguantan emojis ni símbolos raros; los cambiamos por "?"
    return str(valor if valor is not None else "").encode("latin-1", "replace").decode("latin-1")


@app.route("/reporte_usuarios_pdf")                        # Reporte de usuarios en PDF
def reporte_usuarios_pdf():
    if not es_admin():
        return redirect(url_for("login"))
    lista = consultar(
        "SELECT u.id, u.nombre, u.apellido, u.correo, r.nombre AS rol, u.activo "
        "FROM usuarios u JOIN roles r ON r.id = u.rol_id ORDER BY u.id")

    pdf = FPDF(orientation="L")                                  # Creamos el PDF en horizontal (L = landscape)
    pdf.add_page()                                               # Le agregamos una hoja
    pdf.set_font("Helvetica", "B", 16)                           # Letra grande en negrita para el título
    pdf.cell(0, 10, "Reporte de usuarios - Tienda de Ascensores", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)                            # Letra normal
    pdf.cell(0, 8, f"Generado el {datetime.now():%d/%m/%Y %H:%M}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)                                                    # Un espacio

    columnas = [("ID", 15), ("Nombre", 60), ("Correo", 90), ("Rol", 45), ("Activo", 25)]  # Título y ancho
    pdf.set_font("Helvetica", "B", 10)                           # Encabezados en negrita
    pdf.set_fill_color(220, 230, 250)                            # Fondo azulito para los encabezados
    for titulo, ancho in columnas:
        pdf.cell(ancho, 8, titulo, border=1, fill=True)          # Pintamos cada celda del encabezado
    pdf.ln()                                                     # Bajamos de línea

    pdf.set_font("Helvetica", "", 10)                            # Las filas en letra normal
    for u in lista:                                              # Una fila por usuario
        fila = [u["id"], f'{u["nombre"]} {u["apellido"]}', u["correo"], u["rol"], "Sí" if u["activo"] else "No"]
        for (titulo, ancho), valor in zip(columnas, fila):       # zip junta cada columna con su valor
            pdf.cell(ancho, 8, texto_pdf(valor), border=1)
        pdf.ln()

    pdf.ln(4)
    pdf.cell(0, 8, f"Total de usuarios: {len(lista)}")          # Pie con el total

    archivo = BytesIO(pdf.output())                              # pdf.output() devuelve los bytes del PDF
    nombre = f"usuarios_{datetime.now():%Y%m%d}.pdf"
    return send_file(archivo, as_attachment=True, download_name=nombre, mimetype="application/pdf")


# =====================================================================
# Manejo de errores: si la base falla en cualquier página, mostramos un aviso en vez de una pantalla fea
# =====================================================================
@app.errorhandler(MySQLError)
def error_base_datos(error):
    flash("Hubo un problema con la base de datos. Revisa que MySQL esté encendido.", "error")
    return render_template("index.html"), 500    # Lo mandamos al inicio con el aviso


# Si escriben una dirección que no existe (error 404), los mandamos al inicio con un aviso
@app.errorhandler(404)
def pagina_no_encontrada(error):
    # El navegador pide solito el iconito de la pestaña (favicon) y archivos estáticos;
    # si no existen, respondemos 404 calladitos para no mostrar un aviso que nadie pidió
    if request.path.startswith("/static") or request.path == "/favicon.ico":
        return "", 404
    flash("Esa página no existe. Te trajimos al inicio.", "error")
    return redirect(url_for("index"))


# Si abren a mano una dirección que solo funciona con un botón (error 405), igual los devolvemos
@app.errorhandler(405)
def metodo_no_permitido(error):
    flash("Esa acción solo se puede hacer desde su botón.", "error")
    return redirect(url_for("index"))


# Esto solo corre cuando ejecutamos "python app.py" en el computador
if __name__ == "__main__":
    app.run(debug=True)                          # debug=True recarga solo cuando cambiamos el código
